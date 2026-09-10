# -*- coding: utf-8 -*-
"""
Acceptance test ensuring current active configuration maintains 12-epoch authority.
Prevents regression to stale legacy epoch ceilings.
"""

import json
import math
from pathlib import Path
import docx

REPO_ROOT = Path(__file__).resolve().parent.parent

def test_12epoch_authority_constants():
    from research_agent.experiments.training.stage_a2_trainer import StageA2Trainer
    import inspect
    sig = inspect.signature(StageA2Trainer.__init__)
    assert sig.parameters['max_epochs'].default == 12, f"Expected StageA2Trainer default max_epochs=12, got {sig.parameters['max_epochs'].default}"

def test_source_metrics_12epoch_authority():
    p = REPO_ROOT / 'experiments' / 'evidence' / 'stage-a2' / 'reconciliation' / 'CHAPTER3-SOURCE-METRICS.json'
    assert p.exists(), f"Missing {p}"
    data = json.loads(p.read_text(encoding='utf-8'))
    
    arch = data['model_architecture']
    assert arch['max_epochs'] == 12
    assert arch['steps_per_epoch'] == 573
    assert arch['max_optimizer_steps'] == 6876
    assert arch['warmup_steps'] == 343
    
    # Check table rows: no '/ 20' or '/ 12' in epochs_completed
    for s in data['seeds_table']:
        assert isinstance(s['epochs_completed'], int)
        assert 'stop_reason' in s
        assert s['stop_reason'] in ['CEILING_REACHED', 'EARLY_STOPPING', 'HALTED']
        assert '20' not in str(s['epochs_completed'])

def test_execution_plan_v15_12epoch_authority():
    p = REPO_ROOT / 'experiments' / 'plans' / 'STAGE-A2-FIVE-SEED-EXECUTION-PLAN-V1.5.json'
    assert p.exists(), f"Missing {p}"
    data = json.loads(p.read_text(encoding='utf-8'))
    
    hp = data['training_hyperparameters']
    assert hp['max_epochs'] == 12
    assert hp['warmup_steps'] == 343

def test_reconciliation_12epoch_authority():
    p = REPO_ROOT / 'experiments' / 'evidence' / 'stage-a2' / 'reconciliation' / 'STAGE-A2-FIVE-SEED-RECONCILIATION.json'
    assert p.exists(), f"Missing {p}"
    data = json.loads(p.read_text(encoding='utf-8'))
    
    meta = data['metadata']
    assert meta['final_max_epochs'] == 12
    assert meta['optimizer_steps_per_epoch'] == 573
    assert meta['final_max_optimizer_steps'] == 6876
    assert meta['final_warmup_steps'] == 343
    assert meta['canonical_seeds'] == [999]
    assert set(meta['protocol_deviation_seeds']) == {42, 7}
    assert set(meta['noncanonical_seeds']) == {1337, 2024}

def test_docx_no_stale_epoch_references():
    p = REPO_ROOT / 'Chuyên đề chuyên sâu.docx'
    assert p.exists(), f"Missing {p}"
    doc = docx.Document(str(p))
    
    full_text = []
    for para in doc.paragraphs:
        if para.text.strip():
            full_text.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    full_text.append(cell.text)
    
    doc_content = '\n'.join(full_text).lower()
    for forbidden in ['20 epoch', '11460', '11,460', '/ 20', 'trần 20']:
        assert forbidden not in doc_content, f"Forbidden pattern '{forbidden}' found in Master DOCX!"
