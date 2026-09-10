# -*- coding: utf-8 -*-
"""
Automated Test Suite for Table 3.3 Layout & Formatting Verification.
Verifies:
1. Bookmark-based location of Table 3.3 (BK_TBL_3_003).
2. Structural shape: 7 rows x 12 columns.
3. Style: "Table Grid".
4. Table alignment: CENTER (<w:jc w:val="center"/>).
5. Table layout: FIXED (<w:tblLayout w:type="fixed"/>).
6. Row properties: cantSplit enabled across all rows.
7. Cell vertical alignment: CENTER (<w:vAlign w:val="center"/>) across all cells.
8. Cell paragraph formatting:
   - horizontal alignment: CENTER
   - space_before: 0 pt
   - space_after: 0 pt
   - first_line_indent: 0 cm
9. Non-empty cell content.
10. Exact parity between displayed scientific metrics and CHAPTER3-SOURCE-METRICS.json.
"""

import json
from pathlib import Path
import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
import pytest

DOCX_PATH = Path(r"D:\Research\Chuyên đề chuyên sâu.docx")
SOURCE_METRICS_PATH = Path(r"D:\Research\experiments\evidence\stage-a2\reconciliation\CHAPTER3-SOURCE-METRICS.json")


@pytest.fixture(scope="module")
def table_33_and_metrics():
    assert DOCX_PATH.exists(), f"Master DOCX not found at: {DOCX_PATH}"
    assert SOURCE_METRICS_PATH.exists(), f"Source metrics not found at: {SOURCE_METRICS_PATH}"

    doc = docx.Document(str(DOCX_PATH))
    body = doc._body._element

    # Locate Table 3.3 immediately following bookmark BK_TBL_3_003
    target_tbl = None
    for bm in body.iter(qn("w:bookmarkStart")):
        if bm.get(qn("w:name")) == "BK_TBL_3_003":
            curr = bm.getparent()
            while curr is not None:
                curr = curr.getnext()
                if curr is not None and curr.tag.endswith("tbl"):
                    target_tbl = curr
                    break
            break

    assert target_tbl is not None, "Table 3.3 not found following bookmark BK_TBL_3_003!"

    # Find matching docx Table wrapper
    matched_table = None
    for t in doc.tables:
        if t._tbl == target_tbl:
            matched_table = t
            break

    assert matched_table is not None, "Could not map target table to python-docx Table object!"

    with open(SOURCE_METRICS_PATH, "r", encoding="utf-8") as f:
        source_metrics = json.load(f)

    return matched_table, target_tbl, source_metrics


def test_table33_structure_and_style(table_33_and_metrics):
    table, tbl_elem, _ = table_33_and_metrics
    assert len(table.rows) == 7, f"Expected 7 rows, found {len(table.rows)}"
    assert len(table.columns) == 12, f"Expected 12 columns, found {len(table.columns)}"
    assert table.style.name == "Table Grid", f"Expected style 'Table Grid', found '{table.style.name}'"


def test_table33_alignment_and_fixed_layout(table_33_and_metrics):
    table, tbl_elem, _ = table_33_and_metrics
    tblPr = tbl_elem.tblPr
    assert tblPr is not None, "Table missing tblPr!"

    # Table layout: FIXED
    tblLayout = tblPr.find(qn("w:tblLayout"))
    assert tblLayout is not None, "Table missing <w:tblLayout> element!"
    assert tblLayout.get(qn("w:type")) == "fixed", f"Expected fixed layout, got {tblLayout.get(qn('w:type'))}"

    # Table alignment: CENTER
    jc = tblPr.find(qn("w:jc"))
    assert jc is not None, "Table missing <w:jc> element in tblPr!"
    assert jc.get(qn("w:val")) == "center", f"Expected center table alignment, got {jc.get(qn('w:val'))}"
    assert table.alignment == WD_TABLE_ALIGNMENT.CENTER


def test_table33_row_cant_split(table_33_and_metrics):
    table, _, _ = table_33_and_metrics
    for r_idx, row in enumerate(table.rows):
        trPr = row._tr.get_or_add_trPr()
        cantSplit = trPr.find(qn("w:cantSplit"))
        assert cantSplit is not None, f"Row {r_idx} missing <w:cantSplit/>!"


def test_table33_cell_vertical_and_horizontal_alignment(table_33_and_metrics):
    table, _, _ = table_33_and_metrics
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            # Vertical alignment: CENTER
            tcPr = cell._tc.get_or_add_tcPr()
            vAlign = tcPr.find(qn("w:vAlign"))
            assert vAlign is not None, f"Cell ({r_idx}, {c_idx}) missing <w:vAlign/>!"
            assert vAlign.get(qn("w:val")) == "center", (
                f"Cell ({r_idx}, {c_idx}) vAlign is '{vAlign.get(qn('w:val'))}', expected 'center'"
            )

            # Horizontal alignment & paragraph formatting
            for p_idx, p in enumerate(cell.paragraphs):
                txt = p.text.strip()
                if txt:
                    assert p.alignment == WD_ALIGN_PARAGRAPH.CENTER, (
                        f"Cell ({r_idx}, {c_idx}) paragraph {p_idx} alignment is {p.alignment}, expected CENTER"
                    )
                    assert (p.paragraph_format.first_line_indent or 0) == 0, (
                        f"Cell ({r_idx}, {c_idx}) has first_line_indent {p.paragraph_format.first_line_indent}"
                    )
                    assert (p.paragraph_format.space_before or 0) == 0, (
                        f"Cell ({r_idx}, {c_idx}) has space_before {p.paragraph_format.space_before}"
                    )
                    assert (p.paragraph_format.space_after or 0) == 0, (
                        f"Cell ({r_idx}, {c_idx}) has space_after {p.paragraph_format.space_after}"
                    )


def test_table33_no_empty_cells(table_33_and_metrics):
    table, _, _ = table_33_and_metrics
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            txt = cell.text.strip()
            assert len(txt) > 0, f"Cell ({r_idx}, {c_idx}) is unexpectedly empty!"


def test_table33_source_metrics_parity(table_33_and_metrics):
    table, _, source_metrics = table_33_and_metrics
    st = source_metrics["seeds_table"]
    ag_dev = source_metrics["aggregates"]["protocol_deviation"]

    # Check 5 seeds (rows 1..5)
    for i, s in enumerate(st):
        row = table.rows[i + 1]
        cells = [c.text.strip() for c in row.cells]
        assert cells[0] == s["seed_name"]
        assert cells[1] == s["classification"]
        assert cells[2] == str(s["epochs_completed"])
        assert cells[3] == f"{s['steps']:,}"
        assert cells[4] == s["stop_reason"]
        assert cells[5] == f"{s['train_loss']:.4f}"
        assert cells[6] == str(s["best_epoch"])
        assert cells[7] == f"{s['best_val_loss']:.6f}"
        assert cells[8] == f"{s['final_val_loss']:.6f}"
        assert cells[9] == f"{s['l_rel']:.4f}"
        assert cells[10] == f"{s['l_node']:.4f}"
        assert cells[11] == f"{s['l_time']:.4f}"

    # Check aggregate row (row 6)
    agg_row = table.rows[6]
    agg_cells = [c.text.strip() for c in agg_row.cells]
    assert agg_cells[0] == "TB lệch GT\n(3 seed)"
    assert agg_cells[1] == "PROTOCOL_DEVIATION"
    assert agg_cells[2] == "-"
    assert agg_cells[3] == "-"
    assert agg_cells[4] == "-"
    assert agg_cells[5] == f"{ag_dev['mean_final_train_loss']:.4f}"
    assert agg_cells[6] == "-"
    assert agg_cells[7] == f"{ag_dev['mean_best_val_loss']:.6f}"
    assert agg_cells[8] == f"{ag_dev['mean_final_val_loss']:.6f}"
    assert agg_cells[9] == f"{ag_dev['mean_l_rel']:.4f}"
    assert agg_cells[10] == f"{ag_dev['mean_l_node']:.4f}"
    assert agg_cells[11] == f"{ag_dev['mean_l_time']:.4f}"
